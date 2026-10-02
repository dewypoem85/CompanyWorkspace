"""고유명사·용어 후보 정리기 (용어집은 수정하지 않음).

I2Languages.asset 전체 카테고리를 훑어, 현재 용어집에 없는 고유명사/반복 용어 후보를 뽑고
언어별 번역이 일관되는지 함께 집계합니다. AI 호출 없이 동작합니다.

  python proper_noun_audit.py [--asset <I2Languages.asset>] [--out <결과.xlsx>]

산출물 시트
  개요
  고유명사 후보      이름성 항목(스킬·유물·몬스터·시너지 등) — 용어집 등록 여부, 본문 사용 횟수, 번역 일관성
  같은 이름 다른 번역  같은 한글이 키마다 다르게 번역된 경우
  번역 누락          이름 항목인데 특정 언어 번역이 비어 있거나 한글 그대로인 경우
  반복 용어 후보      본문에 자주 나오지만 용어집에 없는 단어 — 언어별 실제 표기 분포 포함
"""
import argparse
import os
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import paths  # noqa: E402
import main as eng  # noqa: E402
from config_manager import load_config  # noqa: E402
from glossary_audit import LANGS, LANG_COL, clean, ngrams  # noqa: E402

HANGUL = re.compile(r"[가-힣]")
SENTENCE_END = re.compile(r"(다|요|죠|까|네|군|지|라|자)[.!?…]*$|[.!?。！？]$")
NAME_CATS = {"Monster", "Pass", "Trait", "Synergy", "Hunting Event", "Period Event", "Score Event", "DungeonOverload", "Boss Event"}
NAME_KEY = re.compile(r"(^|[^a-z])name([^a-z]|$)|\.key$| key$", re.IGNORECASE)

# 단어 끝에 붙는 조사/어미 (긴 것부터 제거)
JOSA = sorted([
    "에서는", "에서의", "으로는", "으로의", "으로서", "에게서", "에서", "에게", "으로", "까지", "부터", "처럼", "보다", "마다", "이며", "이다",
    "하며", "하고", "하여", "하면", "하는", "하게", "한다", "합니다", "하세요", "해서", "해야", "시킨다", "된다", "되는", "되어", "되며",
    "을", "를", "이", "가", "은", "는", "의", "에", "로", "와", "과", "도", "만", "며", "고", "인", "인한",
], key=len, reverse=True)
VERBISH = re.compile(r"(하다|한다|된다|한|한다면|되는|하는|있는|없는|입니다|니다|세요|하세요|해라|합니다|하며|하고|하여|해서|하면|시|중|후|때|동안)$")
STOP = {"모든", "해당", "기본", "추가", "일정", "다음", "이전", "현재", "최대", "최소", "다른", "같은", "각각", "또는", "그리고", "하지만", "이상", "이하", "미만", "초과",
        "사용", "획득", "발동", "적용", "증가", "감소", "효과", "대상", "범위", "시간", "횟수", "확률", "수치", "가능", "불가", "필요", "선택", "확인", "취소",
        "보상", "아이템", "스테이지", "레벨", "단계", "개수", "수량", "남은", "보유", "지급", "구매", "교환", "달성"}


def strip_josa(tok: str) -> str:
    for j in JOSA:
        if tok.endswith(j) and len(tok) - len(j) >= 2:
            return tok[: -len(j)]
    return tok


def is_name_like(r) -> bool:
    kor = r["Korean"].strip()
    if not (2 <= len(kor) <= 20) or not HANGUL.search(kor):
        return False
    if SENTENCE_END.search(kor) or "<" in kor or "{" in kor or "%" in kor or "+" in kor:
        return False
    if re.search(r"\d+[%개회번초]", kor):
        return False
    if NAME_KEY.search(r["Keys"]):
        return True
    return r["_cat"] in NAME_CATS and len(kor) <= 14


def classify(rs) -> str:
    """이름 항목의 유형: 시스템 용어(용어집 후보 1순위) / 스킬·유물·시너지 이름 / 캐릭터·몬스터·NPC 이름 / 이벤트·상품"""
    cats = {r["_cat"] for r in rs}
    keys = " ".join(r["Keys"] for r in rs)
    if cats & {"Monster"} or re.search(r"NPC Name|ArenaNPC|MonsterName|BossName|Skin Name", keys):
        return "캐릭터·몬스터·NPC 이름"
    if cats & {"Artifact", "Skill", "Synergy", "Rune"} and not re.search(r"\.Key$| Key$", keys):
        return "스킬·유물·시너지 이름"
    if cats & {"Hunting Event", "Period Event", "Score Event", "Boss Event", "DungeonOverload", "Shop", "Pass", "MainStory", "CharacterMiniPass", "Mission"}:
        return "이벤트·상품·패스 이름"
    return "시스템 용어(스탯·상태·태그·UI)"


def load_rows(asset):
    rows = []
    for c, n in eng.get_local_i2_categories(asset):
        if c.startswith("#") or n <= 0:
            continue
        _h, rs = eng.parse_local_i2languages_asset(asset, c)
        for r in rs:
            r["Korean"] = (r.get("Korean") or "").strip()
            if r["Korean"]:
                r["_cat"] = c
                rows.append(r)
    return rows


def top_forms(texts, other_counter, n_other, lang):
    """texts(해당 용어가 쓰인 항목의 번역)에서 공통 표기 후보 상위 3개 → [(표기, 항목 수)]"""
    cnt = Counter()
    for t in texts:
        cnt.update(ngrams(t, lang))
    n = len(texts)
    stop = {"the", "of", "a", "an", "to", "and", "in", "de", "la", "el", "los", "las", "del", "y", "en", "un", "una", "por", "con", "que", "is", "for", "on", "with", "your", "you", "se", "al", "su", "es"}
    cands = []
    for g, c in cnt.items():
        if c < 3 or c / n < 0.25:
            continue
        if lang not in {"JPN", "CHS", "CHT"} and all(w in stop for w in g.split()):
            continue
        lift = (c / n) / ((other_counter.get(g, 0) + 1) / (n_other + 1))
        if lift < 25:
            continue
        cands.append((c, lift, len(g), g))
    cands.sort(reverse=True)
    picked = []
    for c, lift, ln, g in cands:
        if any(g in p[0] or p[0] in g for p in picked):
            continue
        picked.append((g, c))
        if len(picked) >= 3:
            break
    return picked


def main():
    cfg = load_config()
    ap = argparse.ArgumentParser()
    ap.add_argument("--asset", default=cfg.get("local_i2_asset_path") or "")
    ap.add_argument("--out", default=os.path.join(paths.OUTPUT_DIR, "고유명사_후보_정리.xlsx"))
    args = ap.parse_args()
    if not args.asset or not os.path.isfile(args.asset):
        print(f"[오류] I2Languages.asset 경로를 찾을 수 없습니다: {args.asset!r}")
        return 1

    print("[1/5] 용어집/게임 텍스트 로드…")
    glossary = eng.fetch_glossary_from_google_sheet("", cfg.get("service_account_json_path", ""))
    rows = load_rows(args.asset)
    gloss_keys = set(glossary)
    gloss_norm = {re.sub(r"\s+", "", k) for k in gloss_keys}
    print(f"      용어집 {len(glossary)}개 / 게임 항목 {len(rows)}개")

    names = [r for r in rows if is_name_like(r)]
    name_keys = {id(r) for r in names}
    body_rows = [r for r in rows if id(r) not in name_keys]
    body_ko = [r["Korean"] for r in body_rows]
    print(f"[2/5] 이름성 항목 {len(names)}개 / 본문 항목 {len(body_rows)}개")

    # 같은 한글 → 항목들
    by_ko = defaultdict(list)
    for r in names:
        by_ko[r["Korean"]].append(r)

    print("[3/5] 이름 항목 사용·일관성 집계…")
    cand = []
    conflict = []
    missing = []
    for ko, rs in sorted(by_ko.items()):
        norm = re.sub(r"\s+", "", ko)
        in_gloss = ko in gloss_keys or norm in gloss_norm
        cats = sorted({r["_cat"] for r in rs})
        uses = [r for r in body_rows if ko in r["Korean"]]
        rec = {"한글": ko, "유형": classify(rs), "카테고리": ", ".join(cats), "항목 수": len(rs), "용어집 등록": "등록" if in_gloss else "미등록",
               "본문 사용 항목 수": len(uses), "대표 Keys": " | ".join(r["Keys"] for r in rs[:2])}
        issue_langs = []
        for lang in LANGS:
            col = LANG_COL[lang]
            vals = [clean((r.get(col) or "")).strip() for r in rs]
            distinct = list(dict.fromkeys(v for v in vals if v))
            rec[lang] = " | ".join(distinct) if distinct else ""
            if len(distinct) > 1:
                conflict.append({"한글": ko, "언어": lang, "카테고리": ", ".join(cats),
                                 "번역 종류 수": len(distinct),
                                 "번역(Keys)": " / ".join(f"{(r.get(col) or '').strip()} [{r['Keys']}]" for r in rs[:6])})
                issue_langs.append(lang)
            if not distinct or all(v == ko for v in distinct):
                missing.append({"한글": ko, "언어": lang, "카테고리": ", ".join(cats), "상태": "번역 없음" if not distinct else "한글 그대로", "Keys": rs[0]["Keys"]})
            # 본문에서 이 이름을 가리키는 부분의 번역이 이름 번역과 일치하는가
            if distinct and uses:
                hit = sum(1 for u in uses if any(d.lower() in clean(u.get(col) or "").lower() for d in distinct))
                rec[f"{lang} 본문 일치율"] = hit / len(uses)
            else:
                rec[f"{lang} 본문 일치율"] = None
        rec["이름 번역이 갈리는 언어"] = ", ".join(issue_langs)
        low = [l for l in LANGS if rec[f"{l} 본문 일치율"] is not None and rec[f"{l} 본문 일치율"] < 0.7 and len(uses) >= 3]
        rec["본문 표기 불일치 언어"] = ", ".join(low)
        cand.append(rec)

    print("[4/5] 반복 용어 후보 추출…")
    def plain_ko(t):
        t = re.sub(r"<[^>]+>|\[[^\]]*\]|\{[^}]*\}", " ", t)
        return t

    EXTRA_STOP = {"경우", "만큼", "동안", "이후", "이전", "동시", "매번", "대신", "이내", "정도", "사이", "근처", "마다", "때문", "이상", "이하", "한번", "하나", "두번", "모두", "자신", "상태에서"}
    josa_df = Counter()
    exact_df = Counter()
    for r in body_rows:
        seen_j, seen_e = set(), set()
        for tok in re.findall(r"[가-힣]{2,}", plain_ko(r["Korean"])):
            base = strip_josa(tok)
            if base != tok and len(base) >= 2:      # 조사가 붙은 형태 = 명사 근거
                seen_j.add(base)
            elif base == tok:
                seen_e.add(tok)
        josa_df.update(seen_j)
        exact_df.update(seen_e)
    df = Counter()
    for stem, nj in josa_df.items():
        if nj >= 3 and stem not in STOP and stem not in EXTRA_STOP and not VERBISH.search(stem) and len(stem) <= 7:
            df[stem] = nj + exact_df.get(stem, 0)
    name_ko_norm = {re.sub(r"\s+", "", ko) for ko in by_ko}
    terms = []
    for tok, n in df.most_common():
        if n < 8:
            break
        if tok in gloss_keys or tok in name_ko_norm:
            continue
        if any(tok == g or tok in g or g in tok for g in gloss_norm if len(g) >= 2):
            continue  # 용어집 용어이거나 그 일부/포함어
        terms.append((tok, n))
    terms = terms[:400]

    all_ng = {lang: Counter() for lang in LANGS}
    for r in rows:
        for lang in LANGS:
            all_ng[lang].update(ngrams(r.get(LANG_COL[lang], ""), lang))
    term_recs = []
    for tok, n in terms:
        users = [r for r in body_rows if tok in plain_ko(r["Korean"])]
        rec = {"한글 후보": tok, "본문 사용 항목 수": len(users), "예시 Keys": " | ".join(u["Keys"] for u in users[:2]),
               "예시 문장": re.sub(r"\s+", " ", plain_ko(users[0]["Korean"]))[:50]}
        for lang in LANGS:
            texts = [u.get(LANG_COL[lang], "") or "" for u in users]
            in_term = Counter()
            for t in texts:
                in_term.update(ngrams(t, lang))
            other = Counter(all_ng[lang])
            other.subtract(in_term)
            forms = top_forms(texts, other, max(1, len(rows) - len(users)), lang)
            rec[lang] = " | ".join(f"{g} ×{c}" for g, c in forms)
        term_recs.append(rec)
    term_recs.sort(key=lambda d: -d["본문 사용 항목 수"])

    print("[5/5] 엑셀 저장…")
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(title, recs, widths=None, pct_cols=()):
        ws = wb.create_sheet(title)
        if not recs:
            ws.append(["(해당 없음)"])
            return
        heads = list(recs[0].keys())
        ws.append(heads)
        for c in range(1, len(heads) + 1):
            cell = ws.cell(1, c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="1E293B")
        for rec in recs:
            row = []
            for h in heads:
                v = rec[h]
                row.append(f"{v:.0%}" if (h in pct_cols and isinstance(v, float)) else v)
            ws.append(row)
        for i, h in enumerate(heads, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = (widths or {}).get(h, 16)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    uncovered = [c for c in cand if c["용어집 등록"] == "미등록"]
    ov = wb.create_sheet("개요")
    ov.column_dimensions["A"].width = 34
    ov.column_dimensions["B"].width = 90
    for a, b in [
        ("항목", "값"),
        ("용어집 용어 수", len(glossary)),
        ("게임 텍스트 항목 수", len(rows)),
        ("이름성 항목 수 / 고유한 한글 이름 수", f"{len(names)} / {len(by_ko)}"),
        ("용어집에 없는 이름 후보", len(uncovered)),
        ("같은 이름이 언어별로 다르게 번역된 경우", len(conflict)),
        ("이름 번역 누락(비어 있음/한글 그대로)", len(missing)),
        ("용어집에 없는 반복 용어 후보", len(term_recs)),
        ("", ""),
        ("판정 방법", "이름성 = 키에 Name/Key가 있거나, 몬스터·시너지·패스·특성·이벤트 카테고리의 짧은 명사구. 문장형·수치 포함 항목은 제외"),
        ("본문 일치율", "그 이름이 한글로 언급된 다른 항목에서, 이름 번역이 그대로 쓰였는지의 비율 (70% 미만 + 3건 이상이면 불일치 언어로 표시)"),
        ("반복 용어 후보", "본문 8개 항목 이상에 나오는 한글 단어(조사 제거). 형태소 분석기 없이 휴리스틱으로 추출하므로 일반 단어가 섞여 있을 수 있음"),
        ("주의", "용어집/게임 데이터는 수정하지 않았습니다. 이 목록은 용어집 정비를 위한 후보입니다. 대조 대상: 로컬 I2Languages.asset(영/일/중간/중번/스)"),
    ]:
        ov.append([a, b])
    for c in (1, 2):
        ov.cell(1, c).font = Font(bold=True)

    pct = tuple(f"{l} 본문 일치율" for l in LANGS)
    type_order = {"시스템 용어(스탯·상태·태그·UI)": 0, "캐릭터·몬스터·NPC 이름": 1, "스킬·유물·시너지 이름": 2, "이벤트·상품·패스 이름": 3}
    cand.sort(key=lambda d: (d["용어집 등록"] != "미등록", type_order.get(d["유형"], 9), -(1 if d["본문 표기 불일치 언어"] else 0), -d["본문 사용 항목 수"], d["한글"]))
    sheet("고유명사 후보", cand, {"한글": 20, "유형": 26, "카테고리": 18, "대표 Keys": 36, **{l: 22 for l in LANGS}}, pct)
    sheet("유형별 요약", [{"유형": t, "용어집 미등록 이름 수": sum(1 for c in uncovered if c["유형"] == t),
                          "그중 본문 표기 불일치 있음": sum(1 for c in uncovered if c["유형"] == t and c["본문 표기 불일치 언어"]),
                          "전체 이름 수": sum(1 for c in cand if c["유형"] == t)} for t in type_order], {"유형": 30})
    conflict.sort(key=lambda d: (-d["번역 종류 수"], d["한글"]))
    sheet("같은 이름 다른 번역", conflict, {"번역(Keys)": 90, "한글": 20})
    sheet("번역 누락", missing, {"Keys": 40, "한글": 22})
    sheet("반복 용어 후보", term_recs, {"한글 후보": 14, "예시 문장": 40, "예시 Keys": 36, **{l: 30 for l in LANGS}})
    wb.save(args.out)
    print(f"\n완료: {args.out}")
    print(f"  이름 {len(by_ko)}개(용어집 미등록 {len(uncovered)}) | 이름 다른 번역 {len(conflict)} | 번역 누락 {len(missing)} | 반복 용어 후보 {len(term_recs)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
